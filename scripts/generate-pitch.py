from pathlib import Path
import json, html
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

root=Path(__file__).resolve().parents[1]; docs=root/"docs"
demo=json.loads((docs/"demo-evaluation.json").read_text()); benchmark=json.loads((root/"public/benchmark-report.json").read_text())
evaluation=demo.get("evaluation",demo);demo={"evaluation":evaluation,"stats":{"alerts":evaluation["attack_alerts"]+evaluation["benign_alerts"],"incidents":next(m for m in evaluation["methods"] if m["id"]=="UT")["queue_size"]}}
methods=demo["evaluation"]["methods"]; ut=next(m for m in methods if m["id"]=="UT"); b0=methods[0]
agg=next(a for a in benchmark["aggregates"] if a["count"]==3000); au=next(m for m in agg["methods"] if m["id"]=="UT"); ab=agg["methods"][0]
stress=next(s for s in benchmark["sweeps"] if s["condition"].get("timing")=="stretched" and s["status"]=="ok"); sm=next(m for m in stress["methods"] if m["id"]=="UT")
colors=["#8996ad","#c3a1f6","#efa863","#7eb3d5","#64e0bb"]
plt.rcParams.update({"font.family":"DejaVu Sans","font.size":12,"axes.facecolor":"#101a27","figure.facecolor":"#101a27","text.color":"#dae7f5","axes.labelcolor":"#a6bad4","xtick.color":"#a6bad4","ytick.color":"#a6bad4","axes.edgecolor":"#344359"})
fig,ax=plt.subplots(figsize=(12,5.3))
for m,c in zip(methods,colors): ax.step([x["k"] for x in m["curve"]],[x["recall"]*100 for x in m["curve"]],where="post",color=c,label=m["name"],linewidth=3 if m["id"]=="UT" else 1.8)
ax.set_xlim(0,1000);ax.set_ylim(0,104);ax.set_xlabel("Opened review items");ax.set_ylabel("Attack episodes surfaced (%)");ax.grid(axis="y",color="#26384e",alpha=.7);ax.legend(loc="lower right",facecolor="#172439",edgecolor="#344359",labelcolor="#dae7f5",fontsize=10);ax.set_title("Seed 239 · 3,000 synthetic alerts · 12 episodes",loc="left",pad=18);fig.tight_layout();fig.savefig(docs/"coverage-curve.png",dpi=190);plt.close(fig)
slides=[
 {"tag":"PROBLEM 25 · SECURITY OPERATIONS","title":"3,000 alerts.\nOne analyst.","body":"Utopia SOC\nTransparent grouping. Human decisions. Measured coverage.","note":"Imagine Utopia · Team 239 · Microsoft Innovate 2026"},
 {"tag":"PROBLEM & USERS","title":"The analyst needs an investigation,\nnot another wall of signals.","cards":[("CONNECT","Bring entity and time evidence together across identity, endpoint, and network alerts."),("PRIORITIZE","Expose asset context, behavior progression, and each risk component."),("VERIFY","Measure how many attack episodes become visible for a given review budget.")],"note":"A research prototype; no claim that triage alone prevents a breach."},
 {"tag":"SOLUTION & ARCHITECTURE","title":"Security decisions stay deterministic.","body":"Validate → deduplicate → correlate → ATT&CK map → classify → score → rank\n\nEvery source alert is retained. Every accepted link has a reason.\n\nSimulation truth → evaluation only. No labels in engine inputs.\nFacts packet → verified template / optional model brief.\nAnalyst decision → versioned state + SHA-256 audit chain.","note":"One shared TypeScript engine · React · local SQLite / hosted D1 · FastAPI adapter"},
 {"tag":"CONTROLLED SIMULATION","title":"Raw events → detectors → validated alerts.","cards":[("FICTIONAL ENTERPRISE","120 hosts · 160 users · 72-hour alert batch. Causal state prerequisites and seeded raw events."),("FOUR ATTACK FAMILIES","Phishing / exfiltration · identity compromise · ransomware · command and control."),("DIFFICULTY CONTROLS","Sensor dropout · duplicates · missing entities · shared hubs · concurrency · timing and severity noise.")],"note":"Controlled ATT&CK v17.1 mappings, validated against the official STIX bundle."},
 {"tag":"WORKING ANALYST INTERFACE","title":"A queue that opens into evidence.","image":"dashboard.png","note":"Live workflow: ranked queue → source timeline → graph reasons → score → human disposition"},
 {"tag":"EXPLAINABILITY & HUMAN CONTROL","title":"Open the relationship. Inspect the reason.","image":"investigation.png","note":"Link contributions, original alert IDs, controlled ATT&CK mappings, and reasoned decisions."},
 {"tag":"FAIR EVALUATION","title":"Same inputs. Same labels. Same metric code.","body":"B0  Severity only — one alert per review item\nB1  Host / 30-minute buckets\nB2  Exact typed-entity sets / time buckets\nB3  Contextual per-alert ranking\nUT  Contextual entity graph + deterministic score\n\nEpisode surfaced = at least one of its alerts in an opened item.\nPairwise grouping quality separately measures reconstruction.","note":"Educational baseline implementations. No vendor product benchmark."},
 {"tag":"MEASURED DEMO · SEED 239","title":f"{demo['stats']['alerts']:,} alerts → {demo['stats']['incidents']:,} review items.","image":"coverage-curve.png","note":f"Full episode coverage: Utopia {ut['items_at_100']} opened items; severity-only {b0['items_at_100']}. Synthetic only; not analyst-time savings."},
 {"tag":"HELD-OUT EVIDENCE","title":"Twenty held-out seeds at each scale.","cards":[(f"{au['metrics']['recall_at_25']['mean']*100:.1f}% RECALL @25",f"Utopia mean across held-out 3,000-alert runs. B0 mean: {ab['metrics']['recall_at_25']['mean']*100:.0f}%."),(f"{au['metrics']['items_at_100']['mean']:.1f} ITEMS TO FULL COVERAGE",f"Utopia mean. Severity-only: {ab['metrics']['items_at_100']['mean']:.1f} items. Total Utopia queue: {au['metrics']['queue_size']['mean']:.1f}."),(f"{au['metrics']['pairwise_f1']['mean']*100:.1f}% GROUPING F1",f"Mean benign contamination: {au['metrics']['benign_contamination']['mean']*100:.1f}%. Surfacing is not full reconstruction.")],"note":"60 holdouts · 40 stress runs · 80 ablations · paired bootstrap intervals · no test-seed tuning"},
 {"tag":"OBSERVED FAILURE CASE","title":"Long gaps fragment the chain.","cards":[("80-MINUTE STEP GAPS",f"Seed 801, 3,000 alerts. Frozen entity windows no longer connect many stages."),(f"{sm['recall_at_25']*100:.1f}% RECALL @25",f"Pairwise recall: {sm['pairwise_recall']*100:.1f}%. Full coverage takes {sm['items_at_100']} items."),("NEXT EXPERIMENT","Measure wider windows, contamination, and segmentation on separate tuning data; never hide the trade-off.")],"note":"This is an actual measured stress result, not a hypothetical caveat."},
 {"tag":"IMPLEMENTATION & LIMITS","title":"Complete workflow. Honest boundaries.","body":"Implemented: simulation, ingest, correlation, evidence, scoring, dashboard, decisions, audit, baselines, benchmarks.\n\nLive brief: deterministic template; optional model adapter is in source.\nAIT-ADS check: checksum verified; 3,000 Wazuh records accepted / 400 AMiner rejected. No gold attack accuracy inferred.\nAudit: trusted-head tamper evidence; no external anchoring.\nNo autonomous response, compliance guarantee, or production accuracy claim.","note":"Tests cover reproducibility, label rejection, preservation, concurrency, owner isolation, and tampering."},
 {"tag":"CONTRIBUTION & NEXT STEP","title":"Make the trade-off visible.","body":"Can transparent grouping surface attack episodes with fewer opened review items?\n\nSynthetic evidence: improved episode visibility, with declining recall and modest grouping F1 as scale increases.\n\nNext: independently labeled external data, a real analyst study, and measured review effort.","note":"Imagine Utopia · Team 239 · Source, benchmark artifacts, and demo supplied"}
]

slides.insert(6,{"tag":"SOC IMPORT WORKBENCH","title":"Bring reports. Keep the evidence.","image":"import-workbench.png","note":"Browser-only parsing · structured event triage / cited narrative assessment · explicit rejection and severity controls"})
slides.insert(10,{"tag":"RESEARCH EXPLORER","title":"Show uncertainty, not just a winning seed.","image":"benchmark-lab.png","note":"Paired comparisons · seed distributions · stress conditions · ablations · temporal replay misses"})
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
bg=RGBColor.from_string("0B1220");mint=RGBColor.from_string("64E0BB");white=RGBColor.from_string("E6EFF9");muted=RGBColor.from_string("A1B4CE")
def box(slide,x,y,w,h,text,size=20,color=white,bold=False):
 shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=shape.text_frame;tf.word_wrap=True;tf.margin_left=0;tf.margin_right=0;tf.margin_top=0
 for i,line in enumerate(text.split("\n")):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name="Aptos";p.font.size=Pt(size);p.font.color.rgb=color;p.font.bold=bold;p.space_after=Pt(12)
 return shape
for i,s in enumerate(slides):
 slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=bg
 box(slide,.6,.4,12,.35,s["tag"],12,mint,True);box(slide,.6,1,12,1.6 if i==0 else 1.2,s["title"],45 if i==0 else 34,white,True)
 if "image"in s:
  imagepath=docs/s["image"]
  if imagepath.exists():
   iw,ih=Image.open(imagepath).size; scale=min(11.8/iw,4.35/ih); w,h=iw*scale,ih*scale;slide.shapes.add_picture(str(imagepath),Inches((13.333-w)/2),Inches(2.25),width=Inches(w),height=Inches(h))
 elif "cards"in s:
  for j,(head,text) in enumerate(s["cards"]):
   x=.65+j*4.2;shape=slide.shapes.add_shape(1,Inches(x),Inches(2.65),Inches(3.95),Inches(3.15));shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string("142238");shape.line.color.rgb=RGBColor.from_string("354B69");box(slide,x+.22,2.92,3.5,.8,head,20,mint,True);box(slide,x+.22,3.86,3.5,1.8,text,19,muted)
 else:box(slide,.65,2.8 if i==0 else 2.35,12,3.9,s["body"],26 if i==0 else 22,muted)
 box(slide,.65,6.95,11.7,.3,s["note"],11,muted);box(slide,12.4,6.95,.4,.3,str(i+1).zfill(2),11,mint)
prs.save(docs/"UTOPIA_PITCH.pptx")
sections=[]
for i,s in enumerate(slides):
 title=html.escape(s["title"]).replace("\n","<br>");inside=f'<div class="tag">{html.escape(s["tag"])}</div><h1>{title}</h1>'
 if "image"in s:inside+=f'<img src="{s["image"]}" alt="{html.escape(s["title"])}">'
 elif "cards"in s:inside+='<div class="cards">'+''.join(f'<article><h2>{html.escape(h)}</h2><p>{html.escape(t)}</p></article>' for h,t in s["cards"])+"</div>"
 else:inside+='<p class="body">'+html.escape(s["body"]).replace("\n","<br>")+"</p>"
 sections.append(f'<section class="slide">{inside}<footer><span>{html.escape(s["note"])}</span><b>{i+1:02}</b></footer></section>')
(docs/"UTOPIA_PITCH.html").write_text('<!doctype html><html lang="en"><head><meta charset="utf-8"><title>Utopia — Hackathon Pitch</title><style>@page{size:1280px 720px;margin:0}*{box-sizing:border-box}body{margin:0;font-family:Arial,sans-serif;background:#0b1220;color:#e6eff9}.slide{width:1280px;height:720px;padding:45px 60px;position:relative;break-after:page;overflow:hidden}.tag{font-size:13px;color:#64e0bb;letter-spacing:2px;font-weight:bold}h1{font-size:45px;line-height:1.18;margin:35px 0 24px;letter-spacing:-1px}.body{font-size:25px;line-height:1.5;color:#a1b4ce;max-width:1120px}.cards{display:grid;grid-template-columns:repeat(3,1fr);gap:23px;margin-top:65px}article{background:#142238;border:1px solid #354b69;border-radius:10px;padding:25px;min-height:270px}h2{font-size:24px;line-height:1.35;color:#64e0bb;margin:0 0 24px}article p{font-size:22px;line-height:1.5;color:#a1b4ce}img{width:1160px;height:440px;object-fit:contain;object-position:left top;background:#101a27}footer{position:absolute;bottom:26px;left:60px;right:50px;display:flex;justify-content:space-between;font-size:13px;color:#a1b4ce;gap:30px}footer b{color:#64e0bb}</style></head><body>'+''.join(sections)+'</body></html>')
print(f"Generated editable {len(slides)}-slide PPTX and printable HTML pitch")
